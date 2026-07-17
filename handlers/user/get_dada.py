"""Excel export helper (shared by keywords/channels)."""

import os

import structlog
from aiogram import Router
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

router = Router(name=__name__)
logger = structlog.get_logger(__name__)


def create_excel_file(data: list, headers: list, filename: str, sheet_name: str) -> str:
    """
    Создаёт Excel-файл с заданными данными, оформлением и автоматической подгонкой ширины столбцов.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_num, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_num, row_data in enumerate(data, start=2):
        for col_num, cell_value in enumerate(row_data, start=1):
            cell = sheet.cell(row=row_num, column=col_num)
            cell.value = cell_value
            cell.alignment = Alignment(horizontal="left", vertical="center")

    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 50)
        sheet.column_dimensions[column_letter].width = adjusted_width

    exports_dir = "exports"
    os.makedirs(exports_dir, exist_ok=True)
    filepath = os.path.join(exports_dir, filename)
    workbook.save(filepath)
    logger.info(f"Excel файл создан: {filepath}")
    return filepath
